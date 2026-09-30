"""Shared utilities for Report mixins."""
from docx import Document
from docx.shared import Pt, Cm
from docx.text.paragraph import Paragraph

from .._xml import (
    add_bookmark_end, add_bookmark_start, add_seq_field,
)
from ..styles import ALIGN


def check_align(value, where='method'):
    """Raise ValueError if `value` is not a valid alignment name.

    None is allowed and means 'use the resolved subconfig value'.
    """
    if value is None:
        return
    if value not in ALIGN:
        allowed = sorted(k for k in ALIGN if k is not None)
        raise ValueError(
            f"invalid align in {where}(): {value!r}; "
            f"expected one of {allowed}"
        )


def check_parts(parts, where='p'):
    """Raise TypeError on the first part that is not str or Inline."""
    from ..inline import Inline
    for part in parts:
        if not isinstance(part, (str, Inline)):
            raise TypeError(
                f"{where}() accepts str or inline nodes "
                f"(b, i, u, s, sup, sub, color, highlight, f, link, ref); "
                f"got {type(part).__name__}"
            )


def apply_layout(para: Paragraph, resolved) -> None:
    """Apply resolved layout properties to a paragraph.

    Accepts any style object that exposes `align`, `line_spacing`,
    `space_before`, `space_after`, `first_line_indent`, `left_indent`.
    Missing attributes and None values are skipped silently.
    """
    align = getattr(resolved, 'align', None)
    if align is not None:
        para.alignment = ALIGN[align]

    pf = para.paragraph_format

    line_spacing = getattr(resolved, 'line_spacing', None)
    if line_spacing is not None:
        pf.line_spacing = line_spacing

    space_before = getattr(resolved, 'space_before', None)
    if space_before is not None:
        pf.space_before = Pt(space_before)

    space_after = getattr(resolved, 'space_after', None)
    if space_after is not None:
        pf.space_after = Pt(space_after)

    first_line_indent = getattr(resolved, 'first_line_indent', None)
    if first_line_indent is not None:
        pf.first_line_indent = Cm(first_line_indent)

    left_indent = getattr(resolved, 'left_indent', None)
    if left_indent is not None:
        pf.left_indent = Cm(left_indent)


def render_caption(doc: Document, caption_style, style_name: str,
                   caption: str, seq_name: str,
                   bookmark_name: str | None = None,
                   bookmark_id: int | None = None) -> Paragraph:
    """Render a numbered caption using a Word SEQ field.

    The template is split at the single `{n}` placeholder. The text
    before it becomes a leading run, the placeholder becomes a SEQ
    field (optionally wrapped in a bookmark for cross-referencing),
    and the text after becomes a trailing run.

    If the template has no `{n}`, no SEQ field is emitted — the
    caption is treated as a plain label.

    Args:
        doc: target document.
        caption_style: a resolved CaptionStyle (prefix, template, align).
        style_name: paragraph style name to apply.
        caption: caption text, substituted for `{caption}`.
        seq_name: SEQ counter name, e.g. 'Table' or 'Figure'.
        bookmark_name: full bookmark name (with prefix) to wrap the
            SEQ field, or None for no bookmark.
        bookmark_id: unique bookmark id, required when bookmark_name
            is given.
    """
    template = caption_style.template

    if template.count('{n}') > 1:
        raise ValueError(
            "caption template contains more than one '{n}' placeholder; "
            "only one SEQ field per caption is supported"
        )

    fmt_args = {
        'prefix': caption_style.prefix,
        'caption': caption,
    }

    if '{n}' in template:
        before_raw, after_raw = template.split('{n}', 1)
        before = before_raw.format(**fmt_args)
        after = after_raw.format(**fmt_args)
        has_seq = True
    else:
        before = template.format(**fmt_args)
        after = ''
        has_seq = False

    para = doc.add_paragraph(style=style_name)
    if before:
        para.add_run(before)
    if has_seq:
        if bookmark_name is not None:
            add_bookmark_start(para, bookmark_name, bookmark_id)
        add_seq_field(para, seq_name)
        if bookmark_name is not None:
            add_bookmark_end(para, bookmark_id)
    if after:
        para.add_run(after)
    para.alignment = ALIGN[caption_style.align]
    return para