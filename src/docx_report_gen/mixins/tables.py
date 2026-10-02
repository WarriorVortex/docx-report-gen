"""Table mixin with optional captions, widths, alignment and merges."""
from typing import Any, Optional, Sequence, Union

from docx.shared import Cm
from docx.table import Table

from .._utils import clear_cell_content
from ..styles import (
    ALIGN, CAPTION_STYLE_NAME, SEQ_TABLE, AlignLiteral,
    resolve_caption, resolve_table,
)
from ._base import DocMixin
from .bookmarks import BookmarkMixin
from ._utils import check_align, render_caption


class TableMixin(BookmarkMixin, DocMixin):
    """Adds table(), merge_cells(), merge_row(), merge_col() to Report."""

    _last_table: Optional[Table]
    _last_table_style: Any

    def __init__(self) -> None:
        super().__init__()
        self._last_table = None
        self._last_table_style = None

    def table(
        self,
        data: Sequence[Sequence[Any]],
        caption: Optional[str] = None,
        caption_align: Optional[AlignLiteral] = None,
        name: Optional[str] = None,
        col_widths: Optional[Union[float, Sequence[float]]] = None,
        col_aligns: Optional[Sequence[AlignLiteral]] = None,
        header_align: Optional[AlignLiteral] = None,
        style: str = 'Table Grid',
        header: bool = True,
    ) -> 'TableMixin':
        """Insert a table, optionally preceded by a numbered caption."""
        if not data:
            raise ValueError('table() requires a non-empty list of rows')
        widths = {len(row) for row in data}
        if len(widths) != 1:
            raise ValueError(
                f'table() rows must have equal length; got {sorted(widths)}'
            )
        check_align(caption_align, where='table caption')
        check_align(header_align, where='table header')

        rows, cols = len(data), len(data[0])
        widths_cm = self._normalize_widths(col_widths, cols)
        aligns = self._normalize_aligns(col_aligns, cols)

        if caption is not None:
            cs = resolve_caption(self.config, {'align': caption_align})
            bookmark_name, bookmark_id = self._prepare_bookmark(
                name, kind='table',
            )
            para = render_caption(
                self.doc, cs, CAPTION_STYLE_NAME,
                caption=caption, seq_name=SEQ_TABLE,
                bookmark_name=bookmark_name, bookmark_id=bookmark_id,
            )
            para.paragraph_format.keep_with_next = True

        resolved = resolve_table(self.config, {
            'col_widths': col_widths,
            'col_aligns': aligns,
            'header_align': header_align,
        })
        t = self.doc.add_table(rows=rows, cols=cols)
        t.style = style

        for i, row in enumerate(data):
            for j, val in enumerate(row):
                cell = t.cell(i, j)
                cell.text = str(val)
                if i == 0 and header:
                    for cp in cell.paragraphs:
                        for run in cp.runs:
                            run.font.bold = True

        if widths_cm is not None:
            for j, width in enumerate(widths_cm):
                for row in t.rows:
                    row.cells[j].width = Cm(width)

        for i in range(rows):
            for j in range(cols):
                align = self._cell_align(
                    i, j, resolved.align, header_align, aligns,
                )
                if align is None:
                    continue
                for cp in t.cell(i, j).paragraphs:
                    cp.alignment = ALIGN[align]

        self._last_table = t
        self._last_table_style = resolved
        return self

    def merge_cells(
        self,
        row1: int,
        col1: int,
        row2: int,
        col2: int,
        text: Optional[str] = None,
    ) -> 'TableMixin':
        """Merge a rectangle of cells in the last-created table."""
        t = self._require_last_table('merge_cells')
        cell_a = t.cell(row1, col1)
        cell_b = t.cell(row2, col2)
        merged = cell_a.merge(cell_b)
        if text is not None:
            self._set_cell_text(merged, text, col1)
        return self

    def merge_row(
        self,
        row: int,
        col1: int,
        col2: int,
        text: Optional[str] = None,
    ) -> 'TableMixin':
        """Merge cells col1..col2 (inclusive) within a single row."""
        return self.merge_cells(row, col1, row, col2, text)

    def merge_col(
        self,
        col: int,
        row1: int,
        row2: int,
        text: Optional[str] = None,
    ) -> 'TableMixin':
        """Merge cells row1..row2 (inclusive) within a single column."""
        return self.merge_cells(row1, col, row2, col, text)

    # ---------- internals ----------

    def _require_last_table(self, where: str) -> Table:
        if self._last_table is None:
            raise RuntimeError(
                f'{where}() requires a preceding table() call'
            )
        return self._last_table

    def _cell_align(
        self,
        i: int,
        j: int,
        default: Optional[AlignLiteral],
        header_align: Optional[AlignLiteral],
        col_aligns: Optional[tuple[AlignLiteral, ...]],
    ) -> Optional[AlignLiteral]:
        if i == 0 and header_align is not None:
            return header_align
        if col_aligns is not None and j < len(col_aligns):
            return col_aligns[j]
        return default

    def _normalize_widths(
        self,
        col_widths: Optional[Union[float, Sequence[float]]],
        cols: int,
    ) -> Optional[tuple[float, ...]]:
        if col_widths is None:
            return None
        if isinstance(col_widths, (int, float)):
            return (float(col_widths),) * cols
        widths = tuple(float(w) for w in col_widths)
        if len(widths) != cols:
            raise ValueError(
                f'col_widths has {len(widths)} entries, '
                f'but the table has {cols} columns'
            )
        return widths

    def _normalize_aligns(
        self,
        col_aligns: Optional[Sequence[AlignLiteral]],
        cols: int,
    ) -> Optional[tuple[AlignLiteral, ...]]:
        if col_aligns is None:
            return None
        aligns = tuple(col_aligns)
        if len(aligns) != cols:
            raise ValueError(
                f'col_aligns has {len(aligns)} entries, '
                f'but the table has {cols} columns'
            )
        for a in aligns:
            check_align(a, where='table col_aligns')
        return aligns

    def _set_cell_text(self, cell: Any, text: str, col: int) -> None:
        """Clear the cell and write fresh text with column alignment."""
        clear_cell_content(cell)
        para = cell.add_paragraph()
        para.add_run(str(text))

        style = self._last_table_style
        align: Optional[AlignLiteral] = None
        if style is not None:
            if style.col_aligns and col < len(style.col_aligns):
                align = style.col_aligns[col]
            else:
                align = style.align
        if align is not None:
            para.alignment = ALIGN[align]