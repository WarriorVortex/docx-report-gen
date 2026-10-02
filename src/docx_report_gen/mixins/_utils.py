"""Shared utilities for Report mixins."""
from typing import Any, Optional, Sequence, Union

from docx.shared import Pt, Cm
from docx.text.paragraph import Paragraph

from ..docx import DocxDocument
from .._utils import (
    add_bookmark_end, add_bookmark_start, add_seq_field,
)
from ..inline import Inline
from ..styles import ALIGN, AlignLiteral, CaptionStyle


def check_align(value: Optional[AlignLiteral],
                where: str = 'method') -> None:
    """Raise ValueError if `value` is not a valid alignment name."""
    if value is None:
        return
    if value not in ALIGN:
        allowed = sorted(k for k in ALIGN if k is not None)
        raise ValueError(
            f"invalid align in {where}(): {value!r}; "
            f"expected one of {allowed}"
        )


def check_parts(parts: Sequence[Union[str, Inline]],
                where: str = 'p') -> None:
    """Raise TypeError on the first part that is not str or Inline."""
    for part in parts:
        if not isinstance(part, (str, Inline)):
            raise TypeError(
                f"{where}() accepts str or inline nodes "
                f"(b, i, u, s, sup, sub, color, highlight, f, link, ref); "
                f"got {type(part).__name__}"
            )


def apply_layout(para: Paragraph, resolved: Any) -> None:
    """Apply resolved layout properties to a paragraph."""
    align = getattr(resolved, 'align', None)
    if align is not None:
        align_value = ALIGN[align]
        if align_value is not None:
            para.alignment = align_value

    pf = para.paragraph_format

    line_spacing = getattr(resolved, 'line_spacing', None)
    if line_spacing is not None:
        pf.line_spacing = line_spacing

    space_before = getattr(resolved, 'space_before', None)
    if space_before is not None:
        pf.space_before = Pt(space_before)

    space_after = getattr(resolved, 'space_after', None)
    if space_after is not None:
        pf.space_after = Pt(space_after)

    first_line_indent = getattr(resolved, 'first_line_indent', None)
    if first_line_indent is not None:
        pf.first_line_indent = Cm(first_line_indent)

    left_indent = getattr(resolved, 'left_indent', None)
    if left_indent is not None:
        pf.left_indent = Cm(left_indent)


def render_caption(
    doc: DocxDocument,
    caption_style: CaptionStyle,
    style_name: str,
    caption: str,
    seq_name: str,
    bookmark_name: Optional[str] = None,
    bookmark_id: Optional[int] = None,
) -> Paragraph:
    """Render a numbered caption using a Word SEQ field."""
    template = caption_style.template

    if template.count('{n}') > 1:
        raise ValueError(
            "caption template contains more than one '{n}' placeholder; "
            "only one SEQ field per caption is supported"
        )

    fmt_args = {
        'prefix': caption_style.prefix,
        'caption': caption,
    }

    if '{n}' in template:
        before_raw, after_raw = template.split('{n}', 1)
        before = before_raw.format(**fmt_args)
        after = after_raw.format(**fmt_args)
        has_seq = True
    else:
        before = template.format(**fmt_args)
        after = ''
        has_seq = False

    para = doc.add_paragraph(style=style_name)
    if before:
        para.add_run(before)
    if has_seq:
        if bookmark_name is not None and bookmark_id is not None:
            add_bookmark_start(para, bookmark_name, bookmark_id)
            add_seq_field(para, seq_name)
            add_bookmark_end(para, bookmark_id)
        else:
            add_seq_field(para, seq_name)
    if after:
        para.add_run(after)

    align_value = ALIGN[caption_style.align]
    if align_value is not None:
        para.alignment = align_value
    return para