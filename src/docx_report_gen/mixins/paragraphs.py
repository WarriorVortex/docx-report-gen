"""Paragraph, block formula and quote mixin."""
from typing import Any, Optional, Union

from docx.enum.text import WD_TAB_ALIGNMENT
from docx.text.paragraph import Paragraph

import math2docx

from ..inline import Inline
from ..styles import (
    QUOTE_STYLE_NAME, AlignLiteral,
    resolve_formula, resolve_paragraph, resolve_quote,
)
from ._base import DocMixin
from .utils import apply_layout, check_align, check_parts


class ParagraphMixin(DocMixin):
    """Adds p(), f() and quote() to Report."""

    _formula_counter: int

    def __init__(self) -> None:
        super().__init__()
        self._formula_counter = 0

    def p(
        self,
        *parts: Union[str, Inline],
        align: Optional[AlignLiteral] = None,
        line_spacing: Optional[float] = None,
        space_before: Optional[int] = None,
        space_after: Optional[int] = None,
        first_line_indent: Optional[float] = None,
    ) -> 'ParagraphMixin':
        """Paragraph from strings and inline nodes."""
        check_align(align, where='p')
        check_parts(parts, where='p')
        local: dict[str, Any] = {
            'align': align,
            'line_spacing': line_spacing,
            'space_before': space_before,
            'space_after': space_after,
            'first_line_indent': first_line_indent,
        }
        resolved = resolve_paragraph(self.config, local)
        para = self.doc.add_paragraph()
        apply_layout(para, resolved)
        for part in parts:
            if isinstance(part, str):
                para.add_run(part)
            else:
                part.render(para)
        return self

    def f(
        self,
        latex: str,
        align: Optional[AlignLiteral] = None,
        number: bool = False,
    ) -> 'ParagraphMixin':
        """Standalone block formula."""
        check_align(align, where='f')
        para = self.doc.add_paragraph()
        if number:
            self._formula_counter += 1
            self._setup_formula_tabs(para)
            para.add_run('\t')
            math2docx.add_math(para, latex)
            para.add_run('\t')
            para.add_run(f'({self._formula_counter})')
        else:
            resolved = resolve_formula(self.config, {'align': align})
            apply_layout(para, resolved)
            math2docx.add_math(para, latex)
        return self

    def quote(
        self,
        *parts: Union[str, Inline],
        align: Optional[AlignLiteral] = None,
        line_spacing: Optional[float] = None,
        space_before: Optional[int] = None,
        space_after: Optional[int] = None,
        first_line_indent: Optional[float] = None,
        left_indent: Optional[float] = None,
    ) -> 'ParagraphMixin':
        """Block quote paragraph."""
        check_align(align, where='quote')
        check_parts(parts, where='quote')
        local: dict[str, Any] = {
            'align': align,
            'line_spacing': line_spacing,
            'space_before': space_before,
            'space_after': space_after,
            'first_line_indent': first_line_indent,
            'left_indent': left_indent,
        }
        resolved = resolve_quote(self.config, local)
        para = self.doc.add_paragraph(style=QUOTE_STYLE_NAME)
        apply_layout(para, resolved)
        for part in parts:
            if isinstance(part, str):
                para.add_run(part)
            else:
                part.render(para)
        return self

    def _setup_formula_tabs(self, para: Paragraph) -> None:
        section = self.doc.sections[0]
        content = (
            section.page_width
            - section.left_margin
            - section.right_margin
        )
        para.paragraph_format.tab_stops.add_tab_stop(
            content / 2, WD_TAB_ALIGNMENT.CENTER,
        )
        para.paragraph_format.tab_stops.add_tab_stop(
            content, WD_TAB_ALIGNMENT.RIGHT,
        )