"""Inline content nodes embedded in Report.p()."""
from typing import Union

import math2docx
from docx.enum.text import WD_COLOR_INDEX
from docx.shared import RGBColor
from docx.text.paragraph import Paragraph

from .._utils import add_hyperlink, add_ref_field
from ..styles.config.types import RGB


HighlightColor = Union[str, WD_COLOR_INDEX]
"""Either a named highlight color or a WD_COLOR_INDEX enum member."""


_HIGHLIGHT_COLORS: dict[str, WD_COLOR_INDEX] = {
    'yellow': WD_COLOR_INDEX.YELLOW,
    'green': WD_COLOR_INDEX.GREEN,
    'blue': WD_COLOR_INDEX.BLUE,
    'red': WD_COLOR_INDEX.RED,
    'gray': WD_COLOR_INDEX.GRAY_25,
}


class Inline:
    """Base class for inline content nodes."""

    def render(self, para: Paragraph) -> None:
        raise NotImplementedError


class Bold(Inline):
    def __init__(self, text: str) -> None:
        self.text: str = text

    def render(self, para: Paragraph) -> None:
        para.add_run(self.text).bold = True


class Italic(Inline):
    def __init__(self, text: str) -> None:
        self.text: str = text

    def render(self, para: Paragraph) -> None:
        para.add_run(self.text).italic = True


class Underline(Inline):
    def __init__(self, text: str) -> None:
        self.text: str = text

    def render(self, para: Paragraph) -> None:
        para.add_run(self.text).underline = True


class Strike(Inline):
    def __init__(self, text: str) -> None:
        self.text: str = text

    def render(self, para: Paragraph) -> None:
        para.add_run(self.text).font.strike = True


class Sup(Inline):
    def __init__(self, text: str) -> None:
        self.text: str = text

    def render(self, para: Paragraph) -> None:
        para.add_run(self.text).font.superscript = True


class Sub(Inline):
    def __init__(self, text: str) -> None:
        self.text: str = text

    def render(self, para: Paragraph) -> None:
        para.add_run(self.text).font.subscript = True


class Color(Inline):
    def __init__(self, text: str, rgb: RGB) -> None:
        self.text: str = text
        self.rgb: RGB = rgb

    def render(self, para: Paragraph) -> None:
        run = para.add_run(self.text)
        run.font.color.rgb = RGBColor(*self.rgb)


class Highlight(Inline):
    def __init__(self,
                 text: str,
                 color: HighlightColor = 'yellow') -> None:
        self.text: str = text
        if isinstance(color, str):
            key = color.lower()
            if key not in _HIGHLIGHT_COLORS:
                raise ValueError(
                    f"unknown highlight color {color!r}; "
                    f"expected one of {sorted(_HIGHLIGHT_COLORS)} "
                    f"or a WD_COLOR_INDEX member"
                )
            self.color: WD_COLOR_INDEX = _HIGHLIGHT_COLORS[key]
        else:
            self.color = color

    def render(self, para: Paragraph) -> None:
        run = para.add_run(self.text)
        run.font.highlight_color = self.color


class Formula(Inline):
    def __init__(self, latex: str) -> None:
        self.latex: str = latex

    def render(self, para: Paragraph) -> None:
        math2docx.add_math(para, self.latex)


class Link(Inline):
    def __init__(self, text: str, url: str) -> None:
        self.text: str = text
        self.url: str = url

    def render(self, para: Paragraph) -> None:
        add_hyperlink(para, self.text, self.url)


class Reference(Inline):
    """Cross-reference to a bookmarked element."""

    def __init__(self, name: str) -> None:
        if not isinstance(name, str) or not name.strip():
            raise ValueError('ref() requires a non-empty name')
        if any(c.isspace() for c in name):
            raise ValueError(
                f'ref() name must not contain whitespace: {name!r}'
            )
        self.name: str = name

    def render(self, para: Paragraph) -> None:
        from ..styles import BOOKMARK_PREFIX
        add_ref_field(para, f'{BOOKMARK_PREFIX}{self.name}')


# ---------- factory functions ----------

def b(text: str) -> Inline:
    """Bold inline node."""
    return Bold(text)


def i(text: str) -> Inline:
    """Italic inline node."""
    return Italic(text)


def u(text: str) -> Inline:
    """Underline inline node."""
    return Underline(text)


def s(text: str) -> Inline:
    """Strikethrough inline node."""
    return Strike(text)


def sup(text: str) -> Inline:
    """Superscript inline node."""
    return Sup(text)


def sub(text: str) -> Inline:
    """Subscript inline node."""
    return Sub(text)


def color(text: str, rgb: RGB) -> Inline:
    """Colored inline node. `rgb` is an (r, g, b) tuple in 0..255."""
    return Color(text, rgb)


def highlight(text: str, color: HighlightColor = 'yellow') -> Inline:
    """Highlighted inline node."""
    return Highlight(text, color)


def f(latex: str) -> Inline:
    """Inline formula node."""
    return Formula(latex)


def link(text: str, url: str) -> Inline:
    """Hyperlink inline node."""
    return Link(text, url)


def ref(name: str) -> Inline:
    """Cross-reference to a bookmark created with table(..., name=...)
    or img(..., name=...). Renders the target's current number.
    """
    return Reference(name)