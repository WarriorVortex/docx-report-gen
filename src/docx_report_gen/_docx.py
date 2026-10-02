"""Structural types for python-docx objects used by the package.

python-docx does not ship a py.typed marker, and its public `Document`
name is a factory function, not a class. Referencing the real class in
annotations is unreliable across mypy setups. Instead, we declare a
narrow Protocol covering exactly the members the package touches.
Real python-docx documents satisfy it structurally — no inheritance
needed, no runtime cost.

Attributes are declared as read-only properties because that is how
python-docx exposes them on the real Document class. Declaring them
as settable variables would make the real Document incompatible with
the protocol.
"""
from typing import Any, Protocol

from docx.table import Table
from docx.text.paragraph import Paragraph


class DocumentProtocol(Protocol):
    """Subset of python-docx's Document used by docx-report-gen."""

    @property
    def styles(self) -> Any: ...

    @property
    def sections(self) -> Any: ...

    @property
    def core_properties(self) -> Any: ...

    @property
    def settings(self) -> Any: ...

    @property
    def paragraphs(self) -> list[Paragraph]: ...

    def add_heading(self, text: str = ...,
                    level: int = ...) -> Paragraph: ...

    def add_paragraph(self, text: str = ...,
                      style: Any = ...) -> Paragraph: ...

    def add_table(self, rows: int, cols: int) -> Table: ...

    def add_page_break(self) -> Paragraph: ...

    def save(self, path: Any) -> None: ...


__all__ = ['DocumentProtocol']