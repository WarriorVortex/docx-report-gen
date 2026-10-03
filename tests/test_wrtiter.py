"""Tests for the imperative free-function API in docx_report_gen.writer.

This file covers the writer module itself: session lifecycle, lazy
creation, delegation of block-level methods, inline factories,
plugin integration and __getattr__ delegation. Runtime configuration
(set_config, apply_styles, set_metadata) lives in test_runtime.py.
"""
import pytest
from docx import Document

from docx_report_gen import writer, Report, StylesConfig, DocumentMetadata
from docx_report_gen.inline import CodeInline, Formula, Inline
from docx_report_gen.plugins import plugins as global_plugins


@pytest.fixture(autouse=True)
def _reset_writer():
    """Drop the singleton session before and after each test."""
    writer.reset()
    yield
    writer.reset()


# ---------- lazy initialization ----------

def test_module_import_does_not_create_report():
    import docx_report_gen.writer as w
    assert w._state._current is None


def test_first_block_call_creates_report():
    import docx_report_gen.writer as w
    writer.h1('Title')
    assert w._state._current is not None


def test_current_creates_report_if_missing():
    import docx_report_gen.writer as w
    r = writer.current()
    assert r is w._state._current


def test_same_report_across_calls():
    writer.h1('A')
    r1 = writer.current()
    writer.p('B')
    r2 = writer.current()
    assert r1 is r2


def test_inline_factory_does_not_create_report():
    """Calling writer.f (an inline factory) is not a block-level call."""
    import docx_report_gen.writer as w
    writer.f('x^2')
    assert w._state._current is None


def test_inline_code_does_not_create_report():
    """Bare code(...) returns an Inline node; does not touch the report."""
    import docx_report_gen.writer as w
    writer.code('print')
    assert w._state._current is None


# ---------- has_session ----------

def test_has_session_false_initially():
    assert writer.has_session() is False


def test_has_session_true_after_call():
    writer.h1('X')
    assert writer.has_session() is True


def test_has_session_false_after_reset():
    writer.h1('X')
    writer.reset()
    assert writer.has_session() is False


# ---------- session: new / configure ----------

def test_new_creates_fresh_report():
    writer.h1('First')
    r1 = writer.current()
    r2 = writer.new()
    assert r1 is not r2
    assert writer.current() is r2


def test_new_returns_report():
    r = writer.new()
    assert isinstance(r, Report)


def test_new_accepts_config():
    cfg = StylesConfig(font='Arial')
    writer.new(config=cfg)
    assert writer.current().config.font == 'Arial'


def test_new_accepts_metadata():
    md = DocumentMetadata(author='Tester')
    writer.new(metadata=md)
    assert writer.current().metadata.author == 'Tester'

# ---------- attach / detach ----------

def test_attach_installs_existing_report():
    r = Report()
    writer.attach(r)
    assert writer.current() is r


def test_attach_returns_same_report():
    r = Report()
    assert writer.attach(r) is r


def test_attach_rejects_non_report():
    with pytest.raises(TypeError, match='Report instance'):
        writer.attach('not a report')


def test_attach_replaces_existing_session():
    writer.h1('Old')
    old = writer.current()
    fresh = Report()
    writer.attach(fresh)
    assert writer.current() is fresh
    assert writer.current() is not old


def test_attach_then_write():
    r = Report()
    writer.attach(r)
    writer.h1('X')
    texts = [p.text for p in r.doc.paragraphs]
    assert 'X' in texts


def test_detach_returns_and_clears():
    writer.h1('X')
    r = writer.current()
    returned = writer.detach()
    assert returned is r
    assert writer.has_session() is False


def test_detach_without_session_returns_none():
    assert writer.detach() is None


def test_detach_then_recreate():
    writer.h1('First')
    r1 = writer.detach()
    writer.h1('Second')
    r2 = writer.current()
    assert r1 is not r2


def test_detach_then_continue_with_report():
    """After detach, the Report remains usable through its own API."""
    writer.h1('X')
    r = writer.detach()
    r.h1('Y')
    texts = [p.text for p in r.doc.paragraphs]
    assert 'X' in texts
    assert 'Y' in texts
    assert writer.has_session() is False


# ---------- reset / close ----------

def test_reset_drops_session():
    import docx_report_gen.writer as w
    writer.h1('X')
    writer.reset()
    assert w._state._current is None


def test_close_drops_session():
    import docx_report_gen.writer as w
    writer.h1('X')
    writer.close()
    assert w._state._current is None


def test_close_idempotent():
    writer.close()
    writer.close()


# ---------- block delegation: headings ----------

def test_h1_writes_paragraph():
    writer.h1('Hello')
    texts = [p.text for p in writer.current().doc.paragraphs]
    assert 'Hello' in texts


def test_all_heading_levels():
    writer.h1('A')
    writer.h2('B')
    writer.h3('C')
    writer.h4('D')
    writer.h5('E')
    writer.h6('F')
    texts = [p.text for p in writer.current().doc.paragraphs]
    assert texts[:6] == ['A', 'B', 'C', 'D', 'E', 'F']


def test_title():
    writer.title('Report')
    texts = [p.text for p in writer.current().doc.paragraphs]
    assert 'Report' in texts


def test_page_break():
    writer.h1('A')
    writer.page_break()
    writer.h1('B')
    texts = [p.text for p in writer.current().doc.paragraphs]
    assert 'A' in texts
    assert 'B' in texts


def test_hr():
    writer.p('above')
    writer.hr()
    writer.p('below')
    texts = [p.text for p in writer.current().doc.paragraphs]
    assert 'above' in texts
    assert 'below' in texts


# ---------- block delegation: content ----------

def test_p_with_strings():
    writer.p('Hello')
    writer.p('World')
    texts = [p.text for p in writer.current().doc.paragraphs]
    assert 'Hello' in texts
    assert 'World' in texts


def test_p_with_inline_nodes():
    writer.p('A: ', writer.b('bold'), ', B: ', writer.i('italic'))
    full = '\n'.join(p.text for p in writer.current().doc.paragraphs)
    assert 'A: bold' in full
    assert 'B: italic' in full


def test_p_with_inline_formula():
    writer.p('See ', writer.f('E = mc^2'))
    paragraphs = writer.current().doc.paragraphs
    assert paragraphs[0].text.startswith('See')


def test_p_with_inline_code():
    writer.p('Function ', writer.code('print'), ' outputs text')
    paragraphs = writer.current().doc.paragraphs
    assert paragraphs[0].text == 'Function print outputs text'


def test_table():
    writer.table([['A', 'B'], [1, 2]])
    assert len(writer.current().doc.tables) == 1


def test_table_with_caption():
    writer.table([['A']], caption='Данные')
    texts = [p.text for p in writer.current().doc.paragraphs]
    assert any('Данные' in t for t in texts)


def test_merge_row_via_writer():
    writer.table([['A', 'B', 'C']])
    writer.merge_row(0, 0, 2, text='merged')
    doc = writer.current().doc
    assert doc.tables[0].cell(0, 0).text == 'merged'


def test_ul_and_ol():
    writer.ul(['a', 'b'])
    writer.ol(['x', 'y'])
    texts = [p.text for p in writer.current().doc.paragraphs]
    assert 'a' in texts
    assert 'x' in texts


def test_quote():
    writer.quote('A quote')
    texts = [p.text for p in writer.current().doc.paragraphs]
    assert 'A quote' in texts


def test_img(png_image):
    writer.img(png_image)
    assert len(writer.current().doc.inline_shapes) == 1


# ---------- accessor: formula ----------

def test_formula_inline_returns_node():
    """Bare f(...) returns an Inline object; does not touch the report."""
    import docx_report_gen.writer as w
    node = writer.f('E = mc^2')
    assert isinstance(node, Formula)
    assert w._state._current is None


def test_formula_block_adds_paragraph():
    writer.f.block('E = mc^2')
    paragraphs = writer.current().doc.paragraphs
    assert len(paragraphs) == 1


def test_formula_block_numbered():
    writer.f.block('x^2', number=True)
    paragraphs = writer.current().doc.paragraphs
    assert '(1)' in paragraphs[0].text


def test_formula_block_unnumbered_does_not_consume_counter():
    writer.f.block('a', number=True)
    writer.f.block('b')
    writer.f.block('c', number=True)
    paragraphs = writer.current().doc.paragraphs
    assert '(1)' in paragraphs[0].text
    assert '(2)' in paragraphs[2].text


def test_formula_block_returns_none():
    assert writer.f.block('x') is None


# ---------- accessor: code ----------

def test_code_inline_returns_node():
    """Bare code(...) returns an Inline object; does not touch the report."""
    import docx_report_gen.writer as w
    node = writer.code('print')
    assert isinstance(node, CodeInline)
    assert w._state._current is None


def test_code_block_adds_paragraphs():
    writer.code.block('x = 1\ny = 2')
    paras = writer.current().doc.paragraphs
    assert len(paras) == 2
    assert paras[0].text == 'x = 1'
    assert paras[1].text == 'y = 2'


def test_code_block_accepts_language():
    writer.code.block('x = 1', language='python')
    paras = writer.current().doc.paragraphs
    assert paras[0].text == 'x = 1'


def test_code_block_accepts_align():
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    writer.code.block('x', align='center')
    paras = writer.current().doc.paragraphs
    assert paras[0].alignment == WD_ALIGN_PARAGRAPH.CENTER


def test_code_block_returns_none():
    assert writer.code.block('x') is None


# ---------- block delegation: page furniture ----------

def test_toc_and_update():
    writer.toc(title='Содержание')
    writer.update_toc()
    texts = [p.text for p in writer.current().doc.paragraphs]
    assert 'Содержание' in texts


def test_header_footer_page_numbers():
    writer.h1('X')
    writer.header('Header text')
    writer.footer('Footer text')
    writer.page_numbers()
    doc = writer.current().doc
    assert any(
        'Header text' in p.text for p in doc.sections[0].header.paragraphs
    )


# ---------- inline factories ----------

def test_inline_factories_available():
    assert writer.b('x') is not None
    assert writer.i('x') is not None
    assert writer.u('x') is not None
    assert writer.s('x') is not None
    assert writer.sup('x') is not None
    assert writer.sub('x') is not None
    assert writer.color('x', (0, 0, 0)) is not None
    assert writer.highlight('x') is not None
    assert writer.f('x^2') is not None
    assert writer.code('x = 1') is not None
    assert writer.link('x', 'https://e.com') is not None
    assert writer.ref('name') is not None


# ---------- lifecycle: save ----------

def test_save_writes_file(tmp_path):
    writer.h1('Title')
    writer.p('Body')
    path = tmp_path / 'out.docx'
    writer.save(str(path))

    assert path.exists()
    doc = Document(str(path))
    texts = [p.text for p in doc.paragraphs]
    assert 'Title' in texts
    assert 'Body' in texts


def test_save_does_not_reset(tmp_path):
    writer.h1('A')
    path1 = tmp_path / 'a.docx'
    writer.save(str(path1))

    writer.h1('B')
    path2 = tmp_path / 'b.docx'
    writer.save(str(path2))

    doc2 = Document(str(path2))
    texts = [p.text for p in doc2.paragraphs]
    assert 'A' in texts
    assert 'B' in texts


# ---------- session independence ----------

def test_two_sessions_isolated(tmp_path):
    writer.new()
    writer.h1('Session 1')
    path1 = tmp_path / 'a.docx'
    writer.save(str(path1))

    writer.new()
    writer.h1('Session 2')
    path2 = tmp_path / 'b.docx'
    writer.save(str(path2))

    doc1 = Document(str(path1))
    doc2 = Document(str(path2))
    assert any(p.text == 'Session 1' for p in doc1.paragraphs)
    assert any(p.text == 'Session 2' for p in doc2.paragraphs)
    assert not any(p.text == 'Session 2' for p in doc1.paragraphs)


def test_attach_mixes_writer_and_report_api():
    r = Report()
    writer.attach(r)
    writer.h1('From writer')
    r.p('From report object')
    texts = [p.text for p in r.doc.paragraphs]
    assert 'From writer' in texts
    assert 'From report object' in texts


# ---------- plugins via writer ----------

def test_register_plugin_global():
    from docx_report_gen import Plugin

    class P(Plugin):
        name = 'test-global'

    writer.register_plugin(P())
    r = Report()
    assert 'test-global' in r.plugins.plugin_names()


def test_register_plugin_does_not_affect_current_session():
    """Global registration does not add the plugin to an already-created
    Report's local registry. It becomes visible only through the parent
    chain, which does not modify the local clone set.
    """
    from docx_report_gen import Plugin

    writer.h1('X')
    existing = writer.current()
    local_before = [p.name for p in existing.plugins]

    class P(Plugin):
        name = 'late-global'

    writer.register_plugin(P())
    local_after = [p.name for p in existing.plugins]

    assert local_before == local_after
    assert 'late-global' not in local_after


def test_unregister_plugin_global():
    from docx_report_gen import Plugin

    class P(Plugin):
        name = 'to-remove'

    p = P()
    writer.register_plugin(p)
    writer.unregister_plugin(p)

    r = Report()
    assert 'to-remove' not in r.plugins.plugin_names()


def test_plugins_returns_local_registry():
    writer.h1('X')
    assert writer.plugins() is writer.current().plugins


def test_block_shorthand():
    writer.h1('X')

    def centered(report, text):
        report.p(text, align='center')

    writer.block('centered', centered)
    assert 'centered' in writer.current().plugins.block_names()


def test_inline_node_shorthand():
    writer.h1('X')

    class Boxed(Inline):
        def __init__(self, text):
            self.text = text
        def render(self, para):
            para.add_run(f'[{self.text}]')

    writer.inline_node('boxed', Boxed)
    assert 'boxed' in writer.current().plugins.inline_names()


def test_local_plugin_attached_to_session():
    from docx_report_gen import Plugin

    class P(Plugin):
        name = 'local'

    writer.h1('X')
    writer.plugins().register(P())
    assert 'local' in writer.current().plugins.plugin_names()


# ---------- plugin delegation via __getattr__ ----------

def test_getattr_delegates_block_method():
    writer.h1('X')

    def centered(report, text):
        report.p(text, align='center')

    writer.block('centered', centered)
    writer.centered('middle')

    texts = [p.text for p in writer.current().doc.paragraphs]
    assert 'middle' in texts


def test_getattr_delegates_inline_factory():
    writer.h1('X')

    class Boxed(Inline):
        def __init__(self, text):
            self.text = text
        def render(self, para):
            para.add_run(f'[{self.text}]')

    writer.inline_node('boxed', Boxed)
    node = writer.boxed('important')
    assert isinstance(node, Boxed)


def test_getattr_unknown_raises():
    writer.h1('X')
    with pytest.raises(AttributeError, match='no such method'):
        writer.definitely_not_defined


def test_getattr_private_raises():
    with pytest.raises(AttributeError):
        writer._private


# ---------- return values ----------

def test_block_methods_return_none():
    assert writer.h1('X') is None
    assert writer.p('Y') is None
    assert writer.table([['A']]) is None

# ---------- use_plugin ----------

def test_use_plugin_by_name():
    from docx_report_gen import Plugin

    class P(Plugin):
        name = 'my-plugin'

    writer.h1('X')                          # creates session
    writer.plugins().register(P())
    p = writer.use_plugin('my-plugin')
    assert p is writer.current().plugins.get_plugin('my-plugin')


def test_use_plugin_by_class():
    from docx_report_gen import Plugin

    class ByClassPlugin(Plugin):
        name = 'by-class'

    writer.h1('X')
    writer.plugins().register(ByClassPlugin())
    p = writer.use_plugin(ByClassPlugin)
    assert p is writer.current().plugins.get_plugin('by-class')


def test_use_plugin_by_instance():
    from docx_report_gen import Plugin

    class ByInstancePlugin(Plugin):
        name = 'by-instance'

    writer.h1('X')
    instance = ByInstancePlugin()
    writer.plugins().register(instance)
    p = writer.use_plugin(instance)
    assert p is writer.current().plugins.get_plugin('by-instance')


def test_use_plugin_methods_use_self_report():
    """Plugin methods access self.report; no argument passing needed."""
    from docx_report_gen import Plugin

    class CenteredPlugin(Plugin):
        name = 'centered'

        def centered(self, text):
            self.report.p(text, align='center')

    writer.h1('X')
    writer.plugins().register(CenteredPlugin())
    writer.use_plugin('centered').centered('hello')

    texts = [p.text for p in writer.current().doc.paragraphs]
    assert 'hello' in texts


def test_use_plugin_finds_global_plugin_clone():
    """A globally registered plugin is cloned per session; use_plugin
    returns the clone, not the template."""
    from docx_report_gen import Plugin

    class GlobalPlugin(Plugin):
        name = 'global-here'

    template = GlobalPlugin()
    writer.register_plugin(template)

    writer.h1('X')                          # new session clones globals
    clone = writer.use_plugin('global-here')
    assert clone is not template
    assert clone is writer.current().plugins.get_plugin('global-here')


def test_use_plugin_missing_raises():
    writer.h1('X')
    with pytest.raises(KeyError):
        writer.use_plugin('definitely-not-registered')


def test_use_plugin_error_lists_available():
    from docx_report_gen import Plugin

    class Known(Plugin):
        name = 'known-plugin'

    writer.h1('X')
    writer.plugins().register(Known())
    with pytest.raises(KeyError, match='known-plugin'):
        writer.use_plugin('missing-plugin')


def test_use_plugin_works_with_block_registration():
    """Both conventions coexist: block registration AND direct methods."""
    from docx_report_gen import Plugin

    class BothPlugin(Plugin):
        name = 'both'

        def setup(self, report):
            # Block registration — for `writer.both_shout(...)`
            self.registry.block('both_shout', self._shout)

        def _shout(self, report, text):
            report.p(text.upper())

        def quiet(self, text):
            # Direct method — self.report is set, no argument needed
            self.report.p(text.lower())

    writer.h1('X')
    writer.plugins().register(BothPlugin())

    writer.both_shout('loud')                          # block API
    writer.use_plugin('both').quiet('QUIET')           # direct API

    texts = [p.text for p in writer.current().doc.paragraphs]
    assert 'LOUD' in texts
    assert 'quiet' in texts