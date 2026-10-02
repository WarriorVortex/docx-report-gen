"""Tests for runtime configuration — set_config, apply_styles,
set_metadata, apply_metadata.

These operations let users change styles and metadata after the
Report has been created. They are the entry point for the writer
API's `set_config` / `set_metadata` wrappers and can also be used
directly on any Report instance.

The file covers both layers:

    Report.set_config / apply_styles / set_metadata / apply_metadata
    writer.set_config / apply_styles / set_metadata / apply_metadata

Any test about session management or block delegation lives in
test_writer.py, not here.
"""
import pytest
from docx import Document
from docx.shared import Pt

from docx_report_gen import (
    writer, Report, StylesConfig, DocumentMetadata, HeadingStyle,
)


@pytest.fixture(autouse=True)
def _reset_writer():
    writer.reset()
    yield
    writer.reset()


# ---------- Report.set_config ----------

def test_report_set_config_replaces_config():
    r = Report()
    new_cfg = StylesConfig(font='Arial')
    r.set_config(new_cfg)
    assert r.config is new_cfg


def test_report_set_config_reapplies_normal_style(docx_path):
    r = Report()          # default font = Times New Roman
    r.h1('X')
    r.set_config(StylesConfig(font='Arial'))
    r.save(str(docx_path))

    doc = Document(str(docx_path))
    assert doc.styles['Normal'].font.name == 'Arial'


def test_report_set_config_reapplies_heading_size(docx_path):
    r = Report()
    r.h1('X')
    r.set_config(StylesConfig(heading_sizes=(30, 24, 20, 18, 16, 14)))
    r.save(str(docx_path))

    doc = Document(str(docx_path))
    assert doc.styles['Heading 1'].font.size == Pt(30)


def test_report_set_config_affects_future_captions(docx_path):
    """set_config changes prefix for captions created after the call.

    Captions already rendered keep their original text: they are
    written into the XML at call time and cannot be re-rendered.
    """
    from docx_report_gen import CaptionStyle

    r = Report()
    r.set_config(StylesConfig(caption=CaptionStyle(prefix='Table')))
    r.table([['A']], caption='Данные')
    r.save(str(docx_path))

    doc = Document(str(docx_path))
    texts = [p.text for p in doc.paragraphs]
    assert any('Table' in t for t in texts)


def test_report_set_config_preserves_local_paragraph_overrides(docx_path):
    """A paragraph with explicit align keeps its align after set_config."""
    r = Report()
    r.h1('Local', align='right')
    r.set_config(StylesConfig(heading=HeadingStyle(align='left')))
    r.save(str(docx_path))

    doc = Document(str(docx_path))
    para = doc.paragraphs[0]
    assert str(para.alignment).endswith('RIGHT') or para.alignment == 2


# ---------- Report.apply_styles ----------

def test_report_apply_styles_after_in_place_change(docx_path):
    """Mutating config in place and calling apply_styles re-applies
    font properties to document styles.

    Alignment is deliberately not set on Word styles: it is applied
    per-paragraph at call time (h1('X', align='center')), so this
    test verifies a font attribute instead.
    """
    r = Report()
    r.h1('X')
    r.config.heading.bold = False
    r.apply_styles()
    r.save(str(docx_path))

    doc = Document(str(docx_path))
    style = doc.styles['Heading 1']
    assert style.font.bold is False


def test_report_apply_styles_is_idempotent(docx_path):
    r = Report()
    r.h1('X')
    r.apply_styles()
    r.apply_styles()
    r.save(str(docx_path))

    doc = Document(str(docx_path))
    assert doc.styles['Normal'].font.name == 'Times New Roman'


def test_report_apply_styles_without_config_change(docx_path):
    """Calling apply_styles with no prior change does not break anything."""
    r = Report()
    r.h1('X')
    r.save(str(docx_path))
    r.apply_styles()
    r.save(str(docx_path))


# ---------- Report.set_metadata ----------

def test_report_set_metadata_applied_at_save(docx_path):
    r = Report()
    r.set_metadata(DocumentMetadata(author='Late Author'))
    r.save(str(docx_path))

    doc = Document(str(docx_path))
    assert doc.core_properties.author == 'Late Author'


def test_report_in_place_metadata_change_applied_at_save(docx_path):
    """Mutating r.metadata in place is picked up at save()."""
    r = Report(metadata=DocumentMetadata(author='First'))
    r.metadata.author = 'Second'
    r.save(str(docx_path))

    doc = Document(str(docx_path))
    assert doc.core_properties.author == 'Second'


def test_report_set_metadata_does_not_write_until_save():
    r = Report(metadata=DocumentMetadata(author='Old'))
    r.set_metadata(DocumentMetadata(author='New'))
    # core_properties still has old value until sync or save
    assert r.doc.core_properties.author == 'Old'


def test_report_set_metadata_overrides_constructor(docx_path):
    r = Report(metadata=DocumentMetadata(author='Constructor'))
    r.set_metadata(DocumentMetadata(author='Runtime'))
    r.save(str(docx_path))

    doc = Document(str(docx_path))
    assert doc.core_properties.author == 'Runtime'


# ---------- Report.apply_metadata ----------

def test_report_apply_metadata_writes_immediately():
    r = Report()
    r.set_metadata(DocumentMetadata(title='Now'))
    r.apply_metadata()
    assert r.doc.core_properties.title == 'Now'


def test_report_apply_metadata_without_set_metadata():
    """apply_metadata with unchanged metadata is safe and idempotent."""
    r = Report(metadata=DocumentMetadata(author='X'))
    r.apply_metadata()
    r.apply_metadata()
    assert r.doc.core_properties.author == 'X'


# ---------- save() applies metadata automatically ----------

def test_save_reapplies_metadata(tmp_path):
    """Two saves pick up metadata changes between them."""
    r = Report(metadata=DocumentMetadata(title='First'))
    p1 = tmp_path / 'a.docx'
    r.save(str(p1))

    r.metadata.title = 'Second'
    p2 = tmp_path / 'b.docx'
    r.save(str(p2))

    doc1 = Document(str(p1))
    doc2 = Document(str(p2))
    assert doc1.core_properties.title == 'First'
    assert doc2.core_properties.title == 'Second'


# ---------- writer wrappers: set_config / apply_styles ----------

def test_writer_set_config_replaces():
    writer.h1('X')
    new_cfg = StylesConfig(font='Georgia')
    writer.set_config(new_cfg)
    assert writer.config() is new_cfg


def test_writer_set_config_reapplies(docx_path):
    writer.h1('X')
    writer.set_config(StylesConfig(font='Arial'))
    writer.save(str(docx_path))

    doc = Document(str(docx_path))
    assert doc.styles['Normal'].font.name == 'Arial'


def test_writer_apply_styles_after_in_place_change():
    """Same as above via the writer API."""
    writer.h1('X')
    writer.config().heading.bold = False
    writer.apply_styles()
    s = writer.style('Heading 1')
    assert s.font.bold is False


# ---------- writer wrappers: metadata ----------

def test_writer_set_metadata(docx_path):
    writer.h1('X')
    writer.set_metadata(DocumentMetadata(author='Writer Author'))
    writer.save(str(docx_path))

    doc = Document(str(docx_path))
    assert doc.core_properties.author == 'Writer Author'


def test_writer_apply_metadata():
    writer.h1('X')
    writer.set_metadata(DocumentMetadata(title='Immediate'))
    writer.apply_metadata()
    assert writer.current().doc.core_properties.title == 'Immediate'


def test_writer_metadata_in_place_change(tmp_path):
    writer.h1('X')
    writer.metadata().author = 'Changed'
    path = tmp_path / 'out.docx'
    writer.save(str(path))

    doc = Document(str(path))
    assert doc.core_properties.author == 'Changed'


def test_writer_metadata_accessor_returns_instance():
    writer.h1('X')
    md = writer.metadata()
    assert isinstance(md, DocumentMetadata)


def test_writer_config_accessor_returns_instance():
    writer.h1('X')
    cfg = writer.config()
    assert isinstance(cfg, StylesConfig)


# ---------- interplay: config + metadata in one session ----------

def test_set_config_and_metadata_together(tmp_path):
    writer.h1('Title')
    writer.set_config(StylesConfig(font='Arial'))
    writer.set_metadata(DocumentMetadata(author='Combo'))
    path = tmp_path / 'combo.docx'
    writer.save(str(path))

    doc = Document(str(path))
    assert doc.styles['Normal'].font.name == 'Arial'
    assert doc.core_properties.author == 'Combo'