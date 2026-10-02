"""Integration test for the plugin workflow end-to-end."""
import pytest
from docx import Document

from docx_report_gen import Report, Plugin, StylesConfig
from docx_report_gen.inline import Inline
from docx_report_gen.plugins import plugins as global_plugins

from ._helpers import all_text


class Boxed(Inline):
    """Inline node used by the test plugin."""

    def __init__(self, text):
        self.text = text

    def render(self, para):
        para.add_run(f'[{self.text}]')


class LayoutPlugin(Plugin):
    """Adds a `boxed` inline node and a `centered` block method,
    and appends a header before saving."""

    name = 'layout-test'

    def setup(self, report):
        self.registry.inline('boxed', Boxed)
        self.registry.block('centered', self._centered)
        self.registry.block('note', self._note)

    def before_save(self, report, path):
        report.header('Отчёт сгенерирован автоматически')

    def after_save(self, report, path):
        self.saved_path = path

    def _centered(self, report, text):
        report.p(text, align='center')

    def _note(self, report, text):
        report.p('Замечание: ', Boxed(text))


def test_plugin_registered_globally_applies_to_report(tmp_path, png_image):
    global_plugins.register(LayoutPlugin())
    r = Report()
    assert 'layout-test' in r.plugins.plugin_names()


def test_plugin_block_method(tmp_path):
    global_plugins.register(LayoutPlugin())
    r = Report()
    r.centered('По центру')

    texts = [p.text for p in r.doc.paragraphs]
    assert 'По центру' in texts


def test_plugin_inline_node(tmp_path):
    global_plugins.register(LayoutPlugin())
    r = Report()
    r.p('Значение: ', r.boxed('важное'), '.')

    texts = [p.text for p in r.doc.paragraphs]
    assert any('Значение: [важное].' in t for t in texts)


def test_plugin_before_save_hook(tmp_path):
    global_plugins.register(LayoutPlugin())
    r = Report()
    r.h1('Section')
    path = tmp_path / 'plugin.docx'
    r.save(str(path))

    doc = Document(str(path))
    header_texts = [p.text for p in doc.sections[0].header.paragraphs]
    assert 'Отчёт сгенерирован автоматически' in header_texts


def test_plugin_after_save_receives_path(tmp_path):
    plugin = LayoutPlugin()
    global_plugins.register(plugin)
    r = Report()
    r.h1('X')
    path = tmp_path / 'after.docx'
    r.save(str(path))

    # The clone that ran after_save is not the template `plugin`,
    # so look up the clone on the report.
    clone = r.plugins.get_plugin('layout-test')
    assert clone.saved_path == str(path)


def test_plugin_note_block_uses_inline_node(tmp_path):
    global_plugins.register(LayoutPlugin())
    r = Report()
    r.note('важное наблюдение')

    texts = [p.text for p in r.doc.paragraphs]
    assert any('Замечание:' in t for t in texts)
    assert any('[важное наблюдение]' in t for t in texts)


def test_plugin_cleanup_on_close(tmp_path):
    global_plugins.register(LayoutPlugin())
    r = Report()
    r.h1('X')
    r.close()

    assert len(r.plugins) == 0


def test_local_plugin_only_affects_one_report(tmp_path):
    class OneOff(Plugin):
        name = 'one-off'
        def setup(self, report):
            self.registry.block('special', self._special)
        def _special(self, report):
            report.p('special content')

    r1 = Report(plugins=[OneOff()])
    r2 = Report()

    r1.special()
    assert 'special' in dir(r1)
    assert 'special' not in dir(r2)


def test_local_plugin_direct_api(tmp_path):
    class MarkerPlugin(Plugin):
        name = 'marker'
        def setup(self, report):
            self.registry.block('mark', self._mark)
        def _mark(self, report, text):
            report.p(f'>>> {text}')

    r = Report(plugins=[MarkerPlugin()])
    r.mark('hello')
    texts = [p.text for p in r.doc.paragraphs]
    assert '>>> hello' in texts