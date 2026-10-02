"""Smoke tests: the package imports, runs, saves and reopens.

These tests do not check fine details. They confirm that every public
method can be called on a fresh Report, that the resulting file is a
valid .docx, and that python-docx can read it back.
"""
import zipfile

import pytest
from docx import Document

from docx_report_gen import (
    Report, StylesConfig, DocumentMetadata,
    HeadingStyle, CaptionStyle,
    b, i, u, s, sup, sub, color, highlight, f, link, ref,
)


# ---------- creation and save ----------

def test_report_can_be_created(report):
    assert report is not None
    assert report.doc is not None


def test_report_saves_and_reopens(saved_docx, reopen):
    assert saved_docx.exists()
    assert saved_docx.stat().st_size > 0
    doc = reopen()
    assert isinstance(doc, Document)


def test_saved_docx_is_valid_zip(saved_docx):
    with zipfile.ZipFile(saved_docx) as zf:
        names = set(zf.namelist())
    assert '[Content_Types].xml' in names
    assert 'word/document.xml' in names
    assert '_rels/.rels' in names


# ---------- content ----------

def test_title_and_headings_round_trip(report, docx_path):
    report.title('Title')
    report.h1('Section 1')
    report.h2('Subsection 1.1')
    report.save(str(docx_path))

    doc = Document(str(docx_path))
    texts = [p.text for p in doc.paragraphs]
    assert 'Title' in texts
    assert 'Section 1' in texts
    assert 'Subsection 1.1' in texts


def test_all_heading_levels(report, docx_path):
    for level in range(1, 7):
        getattr(report, f'h{level}')(f'Heading {level}')
    report.save(str(docx_path))

    doc = Document(str(docx_path))
    texts = [p.text for p in doc.paragraphs]
    for level in range(1, 7):
        assert f'Heading {level}' in texts


def test_paragraph_with_inline_nodes(report, docx_path):
    report.p('plain ', b('bold'), ' ', i('italic'), ' ', f('E = mc^2'))
    report.save(str(docx_path))

    doc = Document(str(docx_path))
    full = '\n'.join(p.text for p in doc.paragraphs)
    assert 'plain' in full
    assert 'bold' in full
    assert 'italic' in full


def test_table_with_caption(report, docx_path):
    report.table(
        [['A', 'B'], [1, 2]],
        caption='Данные',
        name='t1',
    )
    report.save(str(docx_path))

    doc = Document(str(docx_path))
    assert len(doc.tables) == 1
    table = doc.tables[0]
    assert table.cell(0, 0).text == 'A'
    assert table.cell(0, 1).text == 'B'
    assert table.cell(1, 0).text == '1'
    assert table.cell(1, 1).text == '2'


def test_image_insertion(report, png_image, docx_path):
    report.img(png_image, caption='Рисунок', width=2)
    report.save(str(docx_path))

    doc = Document(str(docx_path))
    assert len(doc.inline_shapes) == 1


def test_lists_and_code(report, docx_path):
    report.ul(['first', 'second'])
    report.ol(['one', 'two'])
    report.code('def f(x):\n    return x ** 2\n')
    report.save(str(docx_path))

    doc = Document(str(docx_path))
    texts = [p.text for p in doc.paragraphs]
    assert 'first' in texts
    assert 'second' in texts
    assert 'one' in texts
    assert 'two' in texts
    assert any('def f(x)' in t for t in texts)


def test_quote_and_hr(report, docx_path):
    report.quote('A quoted line.')
    report.hr()
    report.save(str(docx_path))

    doc = Document(str(docx_path))
    texts = [p.text for p in doc.paragraphs]
    assert 'A quoted line.' in texts


def test_block_formula(report, docx_path):
    report.f(r'E = mc^2')
    report.f(r'\int_0^1 x^2 \, dx', number=True)
    report.save(str(docx_path))

    doc = Document(str(docx_path))
    # Formula content lives in OMML; python-docx paragraph text is empty.
    # We only verify the document is saved and reopens cleanly.
    assert isinstance(doc, Document)


# ---------- structure ----------

def test_page_break_and_headers(report, docx_path):
    report.h1('First')
    report.page_break()
    report.h1('Second')
    report.header('Header text')
    report.footer('Footer text')
    report.page_numbers()
    report.save(str(docx_path))

    doc = Document(str(docx_path))
    texts = [p.text for p in doc.paragraphs]
    assert 'First' in texts
    assert 'Second' in texts


def test_toc_and_update(report, docx_path):
    report.toc(title='Содержание')
    report.h1('Chapter')
    report.update_toc()
    report.save(str(docx_path))

    doc = Document(str(docx_path))
    full = '\n'.join(p.text for p in doc.paragraphs)
    assert 'Содержание' in full


# ---------- config and metadata ----------

def test_config_is_applied(docx_path):
    config = StylesConfig(
        font='Arial',
        size=11,
        heading=HeadingStyle(align='left', color=(0, 0, 0)),
        caption=CaptionStyle(prefix='Table', align='left'),
    )
    report = Report(config=config)
    report.h1('With custom config')
    report.save(str(docx_path))

    doc = Document(str(docx_path))
    normal = doc.styles['Normal']
    assert normal.font.name == 'Arial'


def test_metadata_is_written(docx_path):
    metadata = DocumentMetadata(
        author='Test Author',
        title='Test Title',
        subject='Test Subject',
    )
    report = Report(metadata=metadata)
    report.save(str(docx_path))

    doc = Document(str(docx_path))
    assert doc.core_properties.author == 'Test Author'
    assert doc.core_properties.title == 'Test Title'
    assert doc.core_properties.subject == 'Test Subject'


# ---------- inline factories ----------

def test_all_inline_factories_construct():
    """Every inline factory builds an object without errors."""
    assert b('x') is not None
    assert i('x') is not None
    assert u('x') is not None
    assert s('x') is not None
    assert sup('x') is not None
    assert sub('x') is not None
    assert color('x', (255, 0, 0)) is not None
    assert highlight('x') is not None
    assert f('x^2') is not None
    assert link('text', 'https://example.com') is not None
    assert ref('some_name') is not None


def test_ref_requires_name():
    with pytest.raises(ValueError):
        ref('')


def test_ref_rejects_whitespace():
    with pytest.raises(ValueError):
        ref('has space')


# ---------- lifecycle ----------

def test_save_multiple_times(report, tmp_path):
    """A Report can be saved more than once."""
    path1 = tmp_path / 'first.docx'
    path2 = tmp_path / 'second.docx'

    report.h1('Once')
    report.save(str(path1))
    report.h1('Twice')
    report.save(str(path2))

    assert path1.exists()
    assert path2.exists()

    doc1 = Document(str(path1))
    doc2 = Document(str(path2))
    assert any(p.text == 'Once' for p in doc1.paragraphs)
    assert any(p.text == 'Twice' for p in doc2.paragraphs)


def test_close_is_idempotent(report):
    report.close()
    report.close()   # second call must not raise