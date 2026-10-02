"""Tests for plugin lifecycle and attribute delegation in Report.

Uses the global `plugins` registry (cleared by conftest between tests)
and per-Report registries to exercise setup, save and close hooks, plus
attribute delegation through __getattr__.
"""
import pytest
from docx import Document

from docx_report_gen import Report, Plugin
from docx_report_gen.inline import Inline
from docx_report_gen.plugins import plugins as global_plugins


# ---------- Plugin base class ----------

def test_base_plugin_name_is_none():
    assert Plugin.name is None


def test_subclass_gets_default_name():
    class MyPlugin(Plugin):
        pass
    assert MyPlugin.name == 'MyPlugin'


def test_subclass_name_not_overridden_if_set():
    class MyPlugin(Plugin):
        name = 'custom'
    assert MyPlugin.name == 'custom'


def test_plugin_init_sets_none_attributes():
    p = Plugin()
    assert p.report is None
    assert p.registry is None


# ---------- setup on registration ----------

def test_local_register_calls_setup(report):
    class P(Plugin):
        name = 'p'
        def setup(self, report):
            self.got_report = report

    p = P()
    report.plugins.register(p)
    assert p.got_report is report
    assert p.report is report


def test_plugin_passed_to_report_init_is_setup():
    class P(Plugin):
        name = 'p'
        def setup(self, report):
            self.got_report = report

    p = P()
    r = Report(plugins=[p])
    assert p.got_report is r


def test_global_register_does_not_call_setup():
    class P(Plugin):
        name = 'p'
        called = False
        def setup(self, report):
            P.called = True

    global_plugins.register(P())
    assert P.called is False


# ---------- global plugins are cloned per Report ----------

def test_global_plugin_applied_to_new_report():
    class P(Plugin):
        name = 'gp'
        def setup(self, report):
            self.ready = True

    template = P()
    global_plugins.register(template)

    r = Report()
    clone = r.plugins.get_plugin('gp')
    assert clone is not None
    assert clone is not template
    assert clone.ready is True
    assert getattr(template, 'ready', False) is False


def test_global_plugin_cloned_once_per_report():
    class P(Plugin):
        name = 'gp'

    global_plugins.register(P())
    r1 = Report()
    r2 = Report()
    c1 = r1.plugins.get_plugin('gp')
    c2 = r2.plugins.get_plugin('gp')
    assert c1 is not c2


def test_extra_plugins_not_cloned():
    """Plugins passed to Report(plugins=[...]) are used as-is."""
    class P(Plugin):
        name = 'p'

    p = P()
    r = Report(plugins=[p])
    assert r.plugins.get_plugin('p') is p


# ---------- save hooks ----------

def test_before_and_after_save_called(report, tmp_path):
    events = []

    class P(Plugin):
        name = 'p'
        def before_save(self, r, path):
            events.append(('before', path))
        def after_save(self, r, path):
            events.append(('after', path))

    report.plugins.register(P())
    path = tmp_path / 'out.docx'
    report.save(str(path))

    assert events == [
        ('before', str(path)),
        ('after', str(path)),
    ]


def test_save_hooks_order(report, tmp_path):
    events = []

    class P1(Plugin):
        name = 'p1'
        def before_save(self, r, path):
            events.append('p1:before')
        def after_save(self, r, path):
            events.append('p1:after')

    class P2(Plugin):
        name = 'p2'
        def before_save(self, r, path):
            events.append('p2:before')
        def after_save(self, r, path):
            events.append('p2:after')

    report.plugins.register(P1())
    report.plugins.register(P2())
    report.save(str(tmp_path / 'out.docx'))

    assert events == [
        'p1:before',
        'p2:before',
        'p2:after',
        'p1:after',
    ]


def test_hook_receives_same_report(report, tmp_path):
    captured = {}

    class P(Plugin):
        name = 'p'
        def before_save(self, r, path):
            captured['r'] = r

    report.plugins.register(P())
    report.save(str(tmp_path / 'out.docx'))
    assert captured['r'] is report


# ---------- close hooks ----------

def test_close_hooks_called(report):
    events = []

    class P(Plugin):
        name = 'p'
        def before_close(self, r):
            events.append('before')
        def after_close(self, r):
            events.append('after')

    report.plugins.register(P())
    report.close()
    assert events == ['before', 'after']


def test_close_hooks_order(report):
    events = []

    class P1(Plugin):
        name = 'p1'
        def before_close(self, r):
            events.append('p1:before')
        def after_close(self, r):
            events.append('p1:after')

    class P2(Plugin):
        name = 'p2'
        def before_close(self, r):
            events.append('p2:before')
        def after_close(self, r):
            events.append('p2:after')

    report.plugins.register(P1())
    report.plugins.register(P2())
    report.close()

    assert events == [
        'p1:before',
        'p2:before',
        'p2:after',
        'p1:after',
    ]


def test_close_clears_local_registry(report):
    class P(Plugin):
        name = 'p'

    report.plugins.register(P())
    assert len(report.plugins) == 1
    report.close()
    assert len(report.plugins) == 0


def test_close_is_idempotent_and_does_not_repeat_hooks(report):
    calls = []

    class P(Plugin):
        name = 'p'
        def before_close(self, r):
            calls.append('before')

    report.plugins.register(P())
    report.close()
    report.close()
    assert calls == ['before']


# ---------- attribute delegation: block ----------

def test_block_method_delegates_to_handler(report):
    class P(Plugin):
        name = 'p'
        def setup(self, r):
            self.registry.block('shout', self._shout)
        def _shout(self, r, text):
            r.p(text.upper())

    report.plugins.register(P())
    report.shout('hello')

    texts = [p.text for p in report.doc.paragraphs]
    assert 'HELLO' in texts


def test_block_handler_receives_report_as_first_arg(report):
    captured = {}

    class P(Plugin):
        name = 'p'
        def setup(self, r):
            self.registry.block('capture', self._capture)
        def _capture(self, r, value):
            captured['report'] = r
            captured['value'] = value

    report.plugins.register(P())
    report.capture(42)
    assert captured['report'] is report
    assert captured['value'] == 42


def test_block_with_kwargs(report):
    captured = {}

    class P(Plugin):
        name = 'p'
        def setup(self, r):
            self.registry.block('kw', self._kw)
        def _kw(self, r, *, a, b=0):
            captured['a'] = a
            captured['b'] = b

    report.plugins.register(P())
    report.kw(a=1, b=2)
    assert captured == {'a': 1, 'b': 2}


# ---------- attribute delegation: inline ----------

def test_inline_node_delegates_to_factory(report, docx_path):
    class Boxed(Inline):
        def __init__(self, text):
            self.text = text
        def render(self, para):
            para.add_run(f'[{self.text}]')

    class P(Plugin):
        name = 'p'
        def setup(self, r):
            self.registry.inline('boxed', Boxed)

    report.plugins.register(P())
    report.p('x: ', report.boxed('a'), ' y')
    report.save(str(docx_path))

    doc = Document(str(docx_path))
    full = '\n'.join(p.text for p in doc.paragraphs)
    assert 'x: [a] y' in full


def test_inline_factory_called_with_args(report):
    captured = {}

    class P(Plugin):
        name = 'p'
        def setup(self, r):
            self.registry.inline('make', self._make)
        def _make(self, text, count=1):
            captured['text'] = text
            captured['count'] = count
            return None

    report.plugins.register(P())
    report.make('a', count=3)
    assert captured == {'text': 'a', 'count': 3}


# ---------- missing attributes ----------

def test_missing_attribute_raises(report):
    with pytest.raises(AttributeError, match='nonexistent_thing'):
        report.nonexistent_thing


def test_private_missing_attribute_raises(report):
    with pytest.raises(AttributeError):
        report._nonexistent_private


def test_missing_attribute_message_mentions_report(report):
    with pytest.raises(AttributeError, match='Report'):
        report.absolutely_not_there


# ---------- dir() ----------

def test_dir_includes_registered_block(report):
    class P(Plugin):
        name = 'p'
        def setup(self, r):
            self.registry.block('my_block', lambda r: None)

    report.plugins.register(P())
    assert 'my_block' in dir(report)


def test_dir_includes_registered_inline(report):
    class P(Plugin):
        name = 'p'
        def setup(self, r):
            self.registry.inline('my_node', lambda: None)

    report.plugins.register(P())
    assert 'my_node' in dir(report)


def test_dir_includes_global_block(report):
    """Global registrations are visible through the local registry."""
    global_plugins.block('global_block', lambda r: None)
    assert 'global_block' in dir(report)