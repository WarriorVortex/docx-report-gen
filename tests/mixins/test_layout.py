"""Tests for LayoutMixin — page_break, hr, header, footer, page_numbers."""
import zipfile

from docx import Document
from lxml import etree

from ._helpers import (
    paragraphs_xml, paragraph_text, has_bottom_border, has_page_break,
)


W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'


# ---------- page_break ----------

def test_page_break_emits_br(report, docx_path):
    report.p('before')
    report.page_break()
    report.p('after')
    report.save(str(docx_path))
    assert has_page_break(docx_path)


def test_page_break_returns_self(report):
    assert report.page_break() is report


# ---------- hr ----------

def test_hr_paragraph_has_bottom_border(report, docx_path):
    report.hr()
    report.save(str(docx_path))
    paras = paragraphs_xml(docx_path)
    assert has_bottom_border(paras[0])


def test_hr_returns_self(report):
    assert report.hr() is report


# ---------- header ----------

def test_header_text_appears(report, docx_path):
    report.header('Header text')
    report.save(str(docx_path))

    doc = Document(str(docx_path))
    header = doc.sections[0].header
    texts = [p.text for p in header.paragraphs]
    assert 'Header text' in texts


def test_header_align_center(report, docx_path):
    report.header('X', align='center')
    report.save(str(docx_path))

    doc = Document(str(docx_path))
    para = doc.sections[0].header.paragraphs[0]
    assert para.alignment is not None


def test_header_returns_self(report):
    assert report.header('X') is report


# ---------- footer ----------

def test_footer_text_appears(report, docx_path):
    report.footer('Footer text')
    report.save(str(docx_path))

    doc = Document(str(docx_path))
    texts = [p.text for p in doc.sections[0].footer.paragraphs]
    assert 'Footer text' in texts


def test_footer_without_text_clears(report, docx_path):
    report.footer()
    report.save(str(docx_path))

    doc = Document(str(docx_path))
    texts = [p.text for p in doc.sections[0].footer.paragraphs]
    assert all(t == '' for t in texts)


# ---------- page_numbers ----------

def test_page_numbers_sets_different_first_page(report, docx_path):
    report.page_numbers(skip_first=True)
    report.save(str(docx_path))

    doc = Document(str(docx_path))
    assert doc.sections[0].different_first_page_header_footer is True


def test_page_numbers_no_skip_first(report, docx_path):
    report.page_numbers(skip_first=False)
    report.save(str(docx_path))

    doc = Document(str(docx_path))
    assert doc.sections[0].different_first_page_header_footer is False


def test_page_numbers_inserts_page_field(report, docx_path):
    report.page_numbers()
    report.save(str(docx_path))

    with zipfile.ZipFile(docx_path) as zf:
        footer_names = [
            n for n in zf.namelist()
            if n.startswith('word/footer') and n.endswith('.xml')
        ]
    assert footer_names, 'no footer part was written'

    found = False
    with zipfile.ZipFile(docx_path) as zf:
        for name in footer_names:
            xml = etree.fromstring(zf.read(name))
            for instr in xml.iter(f'{{{W}}}instrText'):
                if 'PAGE' in (instr.text or ''):
                    found = True
    assert found, 'no PAGE field in footer'


def test_page_numbers_returns_self(report):
    assert report.page_numbers() is report