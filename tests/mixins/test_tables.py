"""Tests for TableMixin — data, captions, widths, alignment, merges."""
import pytest

from docx_report_gen import Report, StylesConfig, TableStyle
from docx_report_gen.docx import to_twips

from ._helpers import (
    paragraphs_xml, paragraph_text, tables_xml, instr_texts,
    bookmark_names, row_grid_spans, cell_widths_cm_twips,
)


# ---------- basics ----------

def test_table_creates_one_table(report, docx_path):
    report.table([['A', 'B'], [1, 2]])
    report.save(str(docx_path))
    tables = tables_xml(docx_path)
    assert len(tables) == 1


def test_table_cells_contain_data(report, docx_path):
    from docx import Document
    report.table([['A', 'B'], [1, 2]])
    report.save(str(docx_path))

    doc = Document(str(docx_path))
    t = doc.tables[0]
    assert t.cell(0, 0).text == 'A'
    assert t.cell(0, 1).text == 'B'
    assert t.cell(1, 0).text == '1'
    assert t.cell(1, 1).text == '2'


def test_table_returns_self(report):
    assert report.table([['A']]) is report


def test_multiple_tables(report, docx_path):
    report.table([['A']])
    report.table([['B']])
    report.save(str(docx_path))
    tables = tables_xml(docx_path)
    assert len(tables) == 2


# ---------- caption ----------

def test_table_caption_creates_caption_paragraph(report, docx_path):
    report.table([['A']], caption='Данные')
    report.save(str(docx_path))
    paras = paragraphs_xml(docx_path)
    texts = [paragraph_text(p) for p in paras]
    assert any('Данные' in t for t in texts)


def test_table_caption_has_seq_field(report, docx_path):
    report.table([['A']], caption='Данные')
    report.save(str(docx_path))
    assert any('SEQ Table' in t for t in instr_texts(docx_path))


def test_table_without_caption_has_no_seq_field(report, docx_path):
    report.table([['A']])
    report.save(str(docx_path))
    assert not any('SEQ Table' in t for t in instr_texts(docx_path))


def test_table_caption_uses_custom_prefix(docx_path):
    config = StylesConfig()
    r = Report(config=config)
    r.table([['A']], caption='Data')
    r.save(str(docx_path))
    paras = paragraphs_xml(docx_path)
    texts = [paragraph_text(p) for p in paras]
    assert any('Таблица' in t for t in texts)


# ---------- bookmark via name ----------

def test_table_name_creates_bookmark(report, docx_path):
    report.table([['A']], caption='Данные', name='results')
    report.save(str(docx_path))
    assert '_Ref_results' in bookmark_names(docx_path)


def test_table_name_without_caption_no_bookmark(report, docx_path):
    """A name is only meaningful with a caption — the label needs a number."""
    report.table([['A']], name='results')
    report.save(str(docx_path))
    assert '_Ref_results' not in bookmark_names(docx_path)


# ---------- col_widths ----------

def test_table_col_widths_sequence(docx_path):
    r = Report()
    r.table([['A', 'B', 'C']], col_widths=(3.0, 4.0, 5.0))
    r.save(str(docx_path))
    widths = cell_widths_cm_twips(docx_path)
    assert widths == [to_twips(3.0), to_twips(4.0), to_twips(5.0)]


def test_table_col_widths_single_number(docx_path):
    r = Report()
    r.table([['A', 'B', 'C']], col_widths=4.0)
    r.save(str(docx_path))
    widths = cell_widths_cm_twips(docx_path)
    assert widths == [to_twips(4.0)] * 3


def test_table_col_widths_wrong_length_raises(report):
    with pytest.raises(ValueError, match='col_widths'):
        report.table([['A', 'B']], col_widths=(1.0, 2.0, 3.0))


# ---------- col_aligns ----------

def test_table_col_aligns(report, docx_path):
    from docx import Document
    report.table(
        [['A', 'B', 'C'], ['x', 'y', 'z']],
        col_aligns=('left', 'center', 'right'),
    )
    report.save(str(docx_path))

    doc = Document(str(docx_path))
    t = doc.tables[0]
    jc_left = t.cell(1, 0).paragraphs[0].alignment
    jc_center = t.cell(1, 1).paragraphs[0].alignment
    jc_right = t.cell(1, 2).paragraphs[0].alignment
    # python-docx maps alignment via WD_ALIGN_PARAGRAPH
    assert str(jc_center).endswith('CENTER') or jc_center == 1
    assert str(jc_right).endswith('RIGHT') or jc_right == 2


def test_table_col_aligns_wrong_length_raises(report):
    with pytest.raises(ValueError, match='col_aligns'):
        report.table([['A', 'B']], col_aligns=('left',))


# ---------- header_align ----------

def test_table_header_align(docx_path):
    from docx import Document
    r = Report()
    r.table(
        [['A', 'B'], ['x', 'y']],
        header_align='center',
    )
    r.save(str(docx_path))

    doc = Document(str(docx_path))
    header_jc = doc.tables[0].cell(0, 0).paragraphs[0].alignment
    assert str(header_jc).endswith('CENTER') or header_jc == 1


# ---------- validation ----------

def test_table_empty_data_raises(report):
    with pytest.raises(ValueError, match='non-empty'):
        report.table([])


def test_table_unequal_rows_raise(report):
    with pytest.raises(ValueError, match='equal length'):
        report.table([['A', 'B'], ['x']])


def test_table_invalid_caption_align_raises(report):
    with pytest.raises(ValueError):
        report.table([['A']], caption='X', caption_align='invalid')


# ---------- merges ----------

def test_merge_row_produces_grid_span(report, docx_path):
    report.table([['A', 'B', 'C'], ['x', 'y', 'z']])
    report.merge_row(0, 0, 2, text='merged')
    report.save(str(docx_path))

    tables = tables_xml(docx_path)
    spans = row_grid_spans(tables[0], 0)
    assert spans == [3]


def test_merge_row_text_appears_in_all_merged_cells(report, docx_path):
    from docx import Document
    report.table([['A', 'B', 'C']])
    report.merge_row(0, 0, 2, text='merged')
    report.save(str(docx_path))

    doc = Document(str(docx_path))
    t = doc.tables[0]
    assert t.cell(0, 0).text == 'merged'
    assert t.cell(0, 1).text == 'merged'
    assert t.cell(0, 2).text == 'merged'


def test_merge_col_shares_text(report, docx_path):
    from docx import Document
    report.table([['A', 'B'], ['x', 'y'], ['p', 'q']])
    report.merge_col(0, 0, 2, text='vertical')
    report.save(str(docx_path))

    doc = Document(str(docx_path))
    t = doc.tables[0]
    assert t.cell(0, 0).text == 'vertical'
    assert t.cell(1, 0).text == 'vertical'
    assert t.cell(2, 0).text == 'vertical'


def test_merge_cells_rectangle(report, docx_path):
    from docx import Document
    report.table([['A', 'B'], ['x', 'y']])
    report.merge_cells(0, 0, 1, 1, text='block')
    report.save(str(docx_path))

    doc = Document(str(docx_path))
    t = doc.tables[0]
    assert t.cell(0, 0).text == 'block'


def test_merge_without_table_raises(report):
    with pytest.raises(RuntimeError, match='table'):
        report.merge_row(0, 0, 1, text='x')


def test_merge_returns_self(report):
    report.table([['A', 'B']])
    assert report.merge_row(0, 0, 1) is report


# ---------- table style from config ----------

def test_table_style_align_default(docx_path):
    from docx import Document
    config = StylesConfig(table=TableStyle(align='center'))
    r = Report(config=config)
    r.table([['A']])
    r.save(str(docx_path))

    doc = Document(str(docx_path))
    cell_jc = doc.tables[0].cell(0, 0).paragraphs[0].alignment
    assert str(cell_jc).endswith('CENTER') or cell_jc == 1